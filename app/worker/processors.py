import asyncio
import random
import re
from pathlib import Path
from typing import Any

from app.models.enums import JobType


def _summarize_text(text: str, max_sentences: int = 3) -> str:
    clean_text = re.sub(r"\s+", " ", text).strip()
    if not clean_text:
        return "No extractable text was found."
    sentences = re.split(r"(?<=[.!?])\s+", clean_text)
    summary = " ".join(sentence for sentence in sentences[:max_sentences] if sentence)
    return summary[:1200]


def _extract_text_from_file(file_path: str) -> tuple[str, int]:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"document not found: {path}")

    if path.suffix.lower() == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF extraction requires pypdf to be installed") from exc

        reader = PdfReader(str(path))
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        return text, len(reader.pages)

    if path.suffix.lower() in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        return text, max(1, text.count("\f") + 1)

    if path.suffix.lower() == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX extraction requires python-docx to be installed") from exc

        document = Document(str(path))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text)
        return text, max(1, len(document.paragraphs))

    raise ValueError(f"unsupported document extension: {path.suffix}")


async def document_ocr(payload: dict[str, Any]) -> dict[str, Any]:
    await asyncio.sleep(float(payload.get("duration", 1.5)))
    if payload.get("file_path"):
        text, pages = await asyncio.to_thread(_extract_text_from_file, str(payload["file_path"]))
        words = re.findall(r"\w+", text)
        return {
            "document_uri": payload.get("file_path"),
            "filename": payload.get("filename"),
            "pages": pages,
            "language": payload.get("language", "en"),
            "text_length": len(text),
            "word_count": len(words),
            "extracted_text_preview": text[:2000],
            "summary": _summarize_text(text) if payload.get("summarize", True) else None,
            "search_indexed": True,
            "artifact": f"documents/ocr/{random.randint(1000, 9999)}.json",
        }

    pages = int(payload.get("pages", random.randint(1, 80)))
    return {
        "document_uri": payload.get("document_uri", "s3://incoming/document.pdf"),
        "pages": pages,
        "language": payload.get("language", "en"),
        "ocr_confidence": round(random.uniform(0.88, 0.99), 3),
        "extracted_entities": payload.get("entities", ["invoice_number", "total", "due_date"]),
        "search_indexed": True,
        "artifact": f"documents/ocr/{random.randint(1000, 9999)}.json",
    }


async def media_transcode(payload: dict[str, Any]) -> dict[str, Any]:
    await asyncio.sleep(float(payload.get("duration", 1.0)))
    return {
        "source_uri": payload.get("source_uri", "s3://incoming/video.mp4"),
        "duration_seconds": float(payload.get("duration_seconds", 95.4)),
        "codec": payload.get("codec", "h264"),
        "container": payload.get("container", "mp4"),
        "renditions": payload.get("renditions", ["1080p", "720p", "480p"]),
        "thumbnail_uri": f"media/thumbs/{random.randint(1000, 9999)}.jpg",
        "manifest_uri": f"media/manifests/{random.randint(1000, 9999)}.m3u8",
    }


async def ai_summarization(payload: dict[str, Any]) -> dict[str, Any]:
    await asyncio.sleep(float(payload.get("duration", 2.0)))
    if payload.get("fail"):
        raise RuntimeError("simulated model provider timeout")
    text = str(payload.get("text", "Customer calls mention latency, onboarding friction, and positive support sentiment."))
    words = re.findall(r"\w+", text)
    return {
        "model": payload.get("model", "gpt-4.1-mini"),
        "input_tokens": max(len(words), 1),
        "summary": " ".join(words[: min(18, len(words))]) or "No text provided.",
        "topics": payload.get("topics", ["latency", "onboarding", "support"]),
        "sentiment": payload.get("sentiment", "mixed_positive"),
    }


async def content_moderation(payload: dict[str, Any]) -> dict[str, Any]:
    await asyncio.sleep(float(payload.get("duration", 0.8)))
    labels = payload.get("labels", {"violence": 0.01, "self_harm": 0.0, "adult": 0.03})
    flagged = any(float(score) >= 0.8 for score in labels.values())
    return {
        "asset_uri": payload.get("asset_uri", "s3://incoming/upload.jpg"),
        "flagged": flagged,
        "labels": labels,
        "action": "review" if flagged else "approve",
        "review_queue": "trust-and-safety" if flagged else None,
    }


PROCESSORS = {
    JobType.DOCUMENT_OCR.value: document_ocr,
    JobType.MEDIA_TRANSCODE.value: media_transcode,
    JobType.AI_SUMMARIZATION.value: ai_summarization,
    JobType.CONTENT_MODERATION.value: content_moderation,
}
